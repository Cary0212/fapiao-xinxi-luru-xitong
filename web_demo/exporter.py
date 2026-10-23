"""Fill the user's DOCX template, extending its repeating blocks when needed."""

from __future__ import annotations

from copy import deepcopy
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _set_cell(cell, value: str) -> None:
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def _replace_in_element(element, old: str, new: str) -> None:
    for text_node in element.iter(qn("w:t")):
        if text_node.text:
            text_node.text = text_node.text.replace(old, new)


def _money(value: str) -> str:
    try:
        return f"{Decimal(str(value)):.2f}"
    except (InvalidOperation, TypeError, ValueError):
        return ""


def _fill_items(table, items: list[dict]) -> None:
    target = max(1, len(items)) + 1
    while len(table.rows) < target:
        new_row = deepcopy(table.rows[-1]._tr)
        table._tbl.append(new_row)
    while len(table.rows) > target:
        table._tbl.remove(table.rows[-1]._tr)
    for index, item in enumerate(items, start=1):
        row = table.rows[index]
        _set_cell(row.cells[0], str(index))
        name = str(item.get("name", "")).strip()
        quantity = str(item.get("quantity", "1")).strip()
        unit = str(item.get("unit", "件")).strip()
        # The template has no quantity column; retain non-unit quantities in the name.
        if quantity not in ("", "1", "1.0", "1.00"):
            name = f"{name}（{quantity}{unit}）"
        _set_cell(row.cells[1], name)
        _set_cell(row.cells[2], _money(item.get("unit_price", "")))
    if not items:
        for cell in table.rows[1].cells:
            _set_cell(cell, "")


def _keep_rows_intact(table, *, link_to_next: bool) -> None:
    for row_index, row in enumerate(table.rows):
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:cantSplit")) is None:
            properties.append(OxmlElement("w:cantSplit"))
        keep_next = link_to_next or row_index < len(table.rows) - 1
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = keep_next


def build_docx(receipts: list[dict], template_path: Path) -> bytes:
    shopping = [receipt for receipt in receipts if receipt.get("kind") == "shopping"]
    taxi = [receipt for receipt in receipts if receipt.get("kind") == "taxi"]
    document = Document(template_path)
    body = document.element.body
    original = list(body)
    if len(document.tables) != 9:
        raise ValueError("Word 模板结构已变化，预期为 9 个表格")

    # Each shopping block is a label paragraph, metadata table, spacer, item table.
    shopping_blocks = [original[3:7], original[7:11], original[11:15], original[15:19]]
    taxi_heading = original[19]
    taxi_block = original[20:22]
    for index in range(min(len(shopping), 4), 4):
        for element in shopping_blocks[index]:
            body.remove(element)
    if not shopping:
        body.remove(original[2])
    for index in range(4, len(shopping)):
        for element in shopping_blocks[-1]:
            clone = deepcopy(element)
            _replace_in_element(clone, "#4", f"#{index + 1}")
            taxi_heading.addprevious(clone)

    if not taxi:
        body.remove(taxi_heading)
        for element in taxi_block:
            body.remove(element)
    for index in range(1, len(taxi)):
        for element in taxi_block:
            clone = deepcopy(element)
            _replace_in_element(clone, "#1", f"#{index + 1}")
            body[-1].addprevious(clone)

    tables = document.tables
    for index, receipt in enumerate(shopping):
        info = tables[index * 2]
        items_table = tables[index * 2 + 1]
        _set_cell(info.cell(1, 1), receipt.get("address", "").strip() or "原票未载明")
        _set_cell(info.cell(2, 1), receipt.get("date", ""))
        _set_cell(info.cell(3, 1), _money(receipt.get("final_amount", "")))
        _fill_items(items_table, receipt.get("items", []))

    for index, receipt in enumerate(taxi):
        info = tables[len(shopping) * 2 + index]
        _set_cell(info.cell(1, 1), str(receipt.get("number", "")))
        _set_cell(info.cell(2, 1), receipt.get("date", ""))
        _set_cell(info.cell(3, 1), _money(receipt.get("final_amount", "")))

    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith("小票 #") or not paragraph.text.strip():
            paragraph.paragraph_format.keep_with_next = True
    for index in range(len(shopping)):
        _keep_rows_intact(tables[index * 2], link_to_next=True)
        _keep_rows_intact(tables[index * 2 + 1], link_to_next=False)
    for index in range(len(taxi)):
        _keep_rows_intact(tables[len(shopping) * 2 + index], link_to_next=False)

    data = BytesIO()
    document.save(data)
    return data.getvalue()
